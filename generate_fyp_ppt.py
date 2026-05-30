"""Generate FYP proposal presentation as a single .pptx file."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Palette — professional, warm, no pink
NAVY = RGBColor(0x1A, 0x2B, 0x4A)
TEAL = RGBColor(0x0D, 0x73, 0x77)
TEAL_LIGHT = RGBColor(0x14, 0xA0, 0x9E)
LAVENDER = RGBColor(0xE8, 0xE4, 0xF3)
CREAM = RGBColor(0xFA, 0xF8, 0xF5)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x4A, 0x55, 0x68)
SOFT_GREEN = RGBColor(0x2D, 0x6A, 0x4F)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.35)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.35), Inches(13.333), Inches(0.08)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"

    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.6), Inches(0.82), Inches(12), Inches(0.45))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = LAVENDER
        sp.font.name = "Segoe UI"


def bullet_box(slide, items, left=0.75, top=1.65, width=11.8, height=5.5, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Segoe UI"
        p.font.color.rgb = SLATE
        p.space_after = Pt(10)
        p.bullet = True


def card(slide, x, y, w, h, title, body, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LAVENDER
    shape.line.width = Pt(1)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.12))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()

    t = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.25), Inches(w - 0.4), Inches(0.5))
    tp = t.text_frame.paragraphs[0]
    tp.text = title
    tp.font.bold = True
    tp.font.size = Pt(16)
    tp.font.color.rgb = NAVY
    tp.font.name = "Segoe UI"

    b = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.75), Inches(w - 0.4), Inches(h - 0.9))
    bf = b.text_frame
    bf.word_wrap = True
    bp = bf.paragraphs[0]
    bp.text = body
    bp.font.size = Pt(12)
    bp.font.color.rgb = SLATE
    bp.font.name = "Segoe UI"


# --- Slide 1: Title ---
s = prs.slides.add_slide(BLANK)
add_bg(s, NAVY)
ribbon = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.2), Inches(13.333), Inches(2.1))
ribbon.fill.solid()
ribbon.fill.fore_color.rgb = TEAL
ribbon.line.fill.background()

title = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.6))
tp = title.text_frame.paragraphs[0]
tp.text = "AI-Driven Adaptive Study Planning\n& Academic Assistance System"
tp.font.size = Pt(40)
tp.font.bold = True
tp.font.color.rgb = WHITE
tp.font.name = "Segoe UI"
tp.alignment = PP_ALIGN.CENTER

sub = s.shapes.add_textbox(Inches(0.8), Inches(3.45), Inches(11.5), Inches(1.2))
sp = sub.text_frame.paragraphs[0]
sp.text = "Final Year Project Proposal"
sp.font.size = Pt(24)
sp.font.color.rgb = WHITE
sp.font.name = "Segoe UI"
sp.alignment = PP_ALIGN.CENTER

meta = s.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.5), Inches(0.8))
mp = meta.text_frame.paragraphs[0]
mp.text = "Presented by: [Your Name]  |  Supervisor: [Supervisor Name]  |  Department of Computer Science"
mp.font.size = Pt(14)
mp.font.color.rgb = LAVENDER
mp.font.name = "Segoe UI"
mp.alignment = PP_ALIGN.CENTER

tag = s.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.5), Inches(0.5))
tagp = tag.text_frame.paragraphs[0]
tagp.text = "Empowering students through intelligent, compassionate academic support"
tagp.font.size = Pt(13)
tagp.font.italic = True
tagp.font.color.rgb = GOLD
tagp.font.name = "Segoe UI"
tagp.alignment = PP_ALIGN.CENTER

# --- Slide 2: Agenda ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Presentation Overview", "What we will cover today")
items = [
    "Introduction & motivation",
    "Problem statement & objectives",
    "Proposed system architecture",
    "Methodology & technology stack",
    "Best-practice case study insights",
    "Timeline & cost plan (2-month modules)",
    "Expected outcomes & conclusion",
]
bullet_box(s, items, top=1.75, size=22)

# --- Slide 3: Introduction ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Introduction", "Why this project matters for students")
bullet_box(
    s,
    [
        "Higher-education students struggle with time management and structured study planning.",
        "Traditional calendars and to-do apps lack intelligence and context about academic workloads.",
        "AI and NLP can interpret assignment documents and generate actionable study insights.",
        "Our platform adapts to each student — planning, executing, and optimizing study workflows.",
        "Goal: transform passive scheduling into an active, personalized academic companion.",
    ],
    top=1.75,
    size=20,
)

# --- Slide 4: Problem ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Problem Statement", "Real challenges students face every semester")
problems = [
    "Cannot break complex assignments into manageable, sequenced tasks",
    "Poor prioritization → missed deadlines and last-minute stress",
    "No clear guidance on how to start or approach new assignments",
    "Unstructured exam preparation without personalized revision plans",
    "No system that learns from behaviour and adapts over time",
    "Existing tools are passive — they store tasks but do not assist thinking",
]
bullet_box(s, problems, top=1.65, size=19)

# --- Slide 5: Objectives ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Project Objectives", "Clear, measurable goals")
objs = [
    "Process and understand assignment documents (PDF, DOC, images)",
    "Build an AI task decomposition and intelligent planning engine",
    "Implement adaptive scheduling from deadlines and user behaviour",
    "Provide an interactive assistant for execution, research, and outlines",
    "Deliver a dedicated Exam Mode with automated study schedules",
    "Evaluate impact on time management, productivity, and student wellbeing",
]
bullet_box(s, objs, top=1.65, size=19)

# --- Slide 6: System overview ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Proposed System", "One unified platform — plan, execute, evaluate")
card(s, 0.6, 1.7, 3.9, 2.3, "Reminder Mode", "Manual tasks & deadlines without AI — quick capture for busy days.", TEAL)
card(s, 4.7, 1.7, 3.9, 2.3, "Assignment Mode", "Upload PDF/DOC/IMG — AI parses, understands, and plans work.", TEAL_LIGHT)
card(s, 8.8, 1.7, 3.9, 2.3, "Exam Mode", "Subjects + dates → personalized schedule, revision cycles, spaced repetition.", SOFT_GREEN)
card(s, 0.6, 4.3, 6.0, 2.5, "Execution Support", "Summaries, research directions, outlines, and interactive Q&A while working.", NAVY)
card(s, 6.9, 4.3, 5.8, 2.5, "Adaptive Learning", "Monitors completion, delays, and study patterns — adjusts priorities and workload.", GOLD)

# --- Slide 7: Pipeline ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Assignment Processing Pipeline", "From upload to daily actionable plan")
steps = [
    "Document Parsing — extract text and metadata from uploaded files",
    "Understanding Layer (AI) — type, topics, difficulty, estimated effort",
    "Task Decomposition — break work into small, achievable sub-tasks",
    "Planning Module — timelines, milestones, and daily workload distribution",
    "Start Assignment Mode — summaries, research hints, outlines, live assistance",
]
bullet_box(s, steps, top=1.65, size=20)

# --- Slide 8: Methodology ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Development Methodology", "Modular & iterative — industry best practice")
phases = [
    "Phase 1 — Requirement analysis: user needs, scope, success metrics",
    "Phase 2 — System design: architecture, data flow, AI pipelines",
    "Phase 3 — Implementation: frontend, backend, database, AI integration",
    "Phase 4 — Integration: connect modules; seamless AI + logic interaction",
    "Phase 5 — Testing & evaluation: functional tests, student pilots, metrics",
]
bullet_box(s, phases, top=1.65, size=19)

# --- Slide 9: Tech stack ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Tools & Technologies", "Modern, scalable, student-friendly stack")
card(s, 0.6, 1.75, 2.9, 1.8, "Frontend", "HTML, CSS, JavaScript, React", TEAL)
card(s, 3.7, 1.75, 2.9, 1.8, "Backend", "Node.js or Django", TEAL)
card(s, 6.8, 1.75, 2.9, 1.8, "Database", "Firebase or PostgreSQL", TEAL_LIGHT)
card(s, 9.9, 1.75, 2.9, 1.8, "AI / NLP", "LLM APIs + document NLP", NAVY)
card(s, 0.6, 3.9, 5.8, 1.8, "Document Processing", "PDF / DOC / image parsing libraries", SOFT_GREEN)
card(s, 6.7, 3.9, 6.1, 1.8, "Deployment", "Cloud hosting, secure auth, responsive UI", GOLD)

# --- Slide 10: Case studies ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Best-Practice Case Study Approach", "Learning from proven solutions — applied thoughtfully")
card(
    s, 0.55, 1.65, 6.1, 2.45,
    "Notion AI + Structured Workspaces",
    "Lesson: combine free-form notes with AI-generated task lists.\nOur application: assignment uploads become structured milestones automatically.",
    TEAL,
)
card(
    s, 6.85, 1.65, 6.0, 2.45,
    "Khan Academy / AI Tutors (e.g. Khanmigo)",
    "Lesson: guided, step-by-step help beats generic answers.\nOur application: Start Assignment Mode with outlines and interactive Q&A.",
    TEAL_LIGHT,
)
card(
    s, 0.55, 4.35, 6.1, 2.45,
    "Anki & Spaced Repetition Systems",
    "Lesson: retention improves with scheduled revision cycles.\nOur application: Exam Mode with revision loops and topic prioritization.",
    SOFT_GREEN,
)
card(
    s, 6.85, 4.35, 6.0, 2.45,
    "Todoist / Adaptive Productivity Apps",
    "Lesson: reschedule based on real completion behaviour.\nOur application: adaptive learning adjusts priorities and workload intensity.",
    NAVY,
)

note = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12), Inches(0.4))
np = note.text_frame.paragraphs[0]
np.text = "Best-case synthesis: proven UX patterns + our unique academic document pipeline"
np.font.size = Pt(12)
np.font.italic = True
np.font.color.rgb = TEAL
np.font.name = "Segoe UI"

# --- Slide 11: Timeline modules ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Project Timeline", "Four 2-month modules — 8 months total")
modules = [
    ("Module 1\nMonths 1–2", "Requirements, UI/UX design,\narchitecture & environment setup"),
    ("Module 2\nMonths 3–4", "Backend APIs, database,\ndocument parsing layer"),
    ("Module 3\nMonths 5–6", "AI understanding, decomposition,\nplanning & Exam Mode"),
    ("Module 4\nMonths 7–8", "Integration, testing,\ndeployment & FYP evaluation"),
]
for i, (title, body) in enumerate(modules):
    x = 0.55 + i * 3.15
    card(s, x, 1.85, 2.95, 3.2, title, body, TEAL if i % 2 == 0 else TEAL_LIGHT)

arrow_note = s.shapes.add_textbox(Inches(0.6), Inches(5.35), Inches(12), Inches(1.2))
ap = arrow_note.text_frame.paragraphs[0]
ap.text = (
    "Each module ends with a review milestone: demo to supervisor, documentation update, "
    "and risk check — ensuring steady, stress-free progress for the whole team."
)
ap.font.size = Pt(15)
ap.font.color.rgb = SLATE
ap.font.name = "Segoe UI"

# --- Slide 12: Cost breakdown ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Project Cost Plan", "Total budget divided across 2-month modules (PKR)")

total_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.95))
total_box.fill.solid()
total_box.fill.fore_color.rgb = TEAL
total_box.line.fill.background()
tt = s.shapes.add_textbox(Inches(0.8), Inches(1.72), Inches(11.8), Inches(0.6))
ttp = tt.text_frame.paragraphs[0]
ttp.text = "TOTAL ESTIMATED PROJECT COST:  PKR 32,000  (≈ PKR 4,000 per month)"
ttp.font.size = Pt(22)
ttp.font.bold = True
ttp.font.color.rgb = WHITE
ttp.font.name = "Segoe UI"
ttp.alignment = PP_ALIGN.CENTER

rows = [
    ("Module 1 (M1–M2)", "Design, tools, initial cloud setup", "6,000", "19%"),
    ("Module 2 (M3–M4)", "Backend, database, parsing libraries", "8,500", "27%"),
    ("Module 3 (M5–M6)", "LLM API usage, AI testing (peak phase)", "11,000", "34%"),
    ("Module 4 (M7–M8)", "Hosting, user testing, deployment, reports", "6,500", "20%"),
]

table_shape = s.shapes.add_table(5, 4, Inches(0.6), Inches(2.75), Inches(12.1), Inches(2.8))
table = table_shape.table
headers = ["Module", "Cost Items", "Amount (PKR)", "% of Total"]
for c, h in enumerate(headers):
    cell = table.cell(0, c)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    for p in cell.text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.font.name = "Segoe UI"

for r, row in enumerate(rows, start=1):
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.text = val
        if r % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = LAVENDER
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.color.rgb = SLATE
            p.font.name = "Segoe UI"

cost_note = s.shapes.add_textbox(Inches(0.6), Inches(5.75), Inches(12), Inches(1.3))
cn = cost_note.text_frame.paragraphs[0]
cn.text = (
    "Includes: cloud hosting, LLM API credits, optional domain, and student pilot incentives. "
    "Excludes personal laptop (already available). Costs are conservative estimates — "
    "free tiers (Firebase, Gemini) may reduce actual spend in best-case deployment."
)
cn.font.size = Pt(12)
cn.font.color.rgb = SLATE
cn.font.name = "Segoe UI"

# --- Slide 13: Cost pie visual (simple bars) ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Cost Distribution by Module", "Where investment goes — transparent planning")
amounts = [6000, 8500, 11000, 6500]
labels = ["M1–M2", "M3–M4", "M5–M6", "M7–M8"]
colors = [TEAL, TEAL_LIGHT, NAVY, SOFT_GREEN]
max_h = 3.0
for i, (amt, lab, col) in enumerate(zip(amounts, labels, colors)):
    h = (amt / 11000) * max_h
    x = 1.2 + i * 2.85
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.2 - h), Inches(1.6), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = col
    bar.line.fill.background()
    lbl = s.shapes.add_textbox(Inches(x), Inches(5.35), Inches(1.6), Inches(0.4))
    lp = lbl.text_frame.paragraphs[0]
    lp.text = lab
    lp.font.size = Pt(14)
    lp.font.bold = True
    lp.font.color.rgb = NAVY
    lp.font.name = "Segoe UI"
    lp.alignment = PP_ALIGN.CENTER
    val = s.shapes.add_textbox(Inches(x), Inches(4.9 - h), Inches(1.6), Inches(0.35))
    vp = val.text_frame.paragraphs[0]
    vp.text = f"PKR {amt:,}"
    vp.font.size = Pt(11)
    vp.font.color.rgb = SLATE
    vp.font.name = "Segoe UI"
    vp.alignment = PP_ALIGN.CENTER

legend = s.shapes.add_textbox(Inches(0.75), Inches(1.75), Inches(12), Inches(2.5))
lf = legend.text_frame
lp0 = lf.paragraphs[0]
lp0.text = "Best-case cost strategy"
lp0.font.bold = True
lp0.font.size = Pt(20)
lp0.font.color.rgb = NAVY
lp0.font.name = "Segoe UI"
for line in [
    "• Use free student / developer tiers for cloud and AI where possible",
    "• Open-source parsing (pdf.js, mammoth) — zero licensing cost",
    "• Phase API spending to Module 3 when features are stable",
    "• Module 4 focuses on polish and evaluation, not heavy new development",
]:
    p = lf.add_paragraph()
    p.text = line
    p.font.size = Pt(16)
    p.font.color.rgb = SLATE
    p.font.name = "Segoe UI"
    p.space_after = Pt(6)

# --- Slide 14: Outcomes ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Expected Outcomes", "Impact we aim to demonstrate")
outcomes = [
    "Improved academic workload management and reduced last-minute stress",
    "Lower procrastination through structured, AI-generated daily plans",
    "Faster, clearer assignment completion with guided execution support",
    "Smarter exam preparation with personalized, adaptive schedules",
    "Measurable gains in productivity and time-management metrics",
    "A reusable, modular codebase suitable for future research extensions",
]
bullet_box(s, outcomes, top=1.65, size=19)

# --- Slide 15: Conclusion ---
s = prs.slides.add_slide(BLANK)
add_bg(s)
add_header_bar(s, "Conclusion", "Bridging traditional tools and intelligent academic assistance")
conc = s.shapes.add_textbox(Inches(0.75), Inches(1.75), Inches(11.8), Inches(3.5))
cf = conc.text_frame
cf.word_wrap = True
p0 = cf.paragraphs[0]
p0.text = (
    "This project addresses a real, everyday problem for university students — not with another "
    "passive calendar, but with an empathetic AI companion that plans, guides, and adapts.\n\n"
    "By combining proven patterns from leading ed-tech and productivity tools with our own "
    "assignment-understanding pipeline, we deliver a practical, research-worthy, and socially "
    "meaningful Final Year Project.\n\n"
    "We respectfully seek your guidance and approval to proceed with development."
)
p0.font.size = Pt(20)
p0.font.color.rgb = SLATE
p0.font.name = "Segoe UI"
p0.line_spacing = 1.3

thanks = s.shapes.add_textbox(Inches(0.75), Inches(5.5), Inches(11.8), Inches(1.2))
tp = thanks.text_frame.paragraphs[0]
tp.text = "Thank you for your time and support."
tp.font.size = Pt(28)
tp.font.bold = True
tp.font.color.rgb = TEAL
tp.font.name = "Segoe UI"
tp.alignment = PP_ALIGN.CENTER

qa = s.shapes.add_textbox(Inches(0.75), Inches(6.2), Inches(11.8), Inches(0.6))
qp = qa.text_frame.paragraphs[0]
qp.text = "Questions & Discussion"
qp.font.size = Pt(18)
qp.font.color.rgb = GOLD
qp.font.name = "Segoe UI"
qp.alignment = PP_ALIGN.CENTER

# Save
out_path = r"c:\Users\owner\OneDrive\Desktop\mean mern\FYP_Proposal_Presentation.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slides: {len(prs.slides)}")
